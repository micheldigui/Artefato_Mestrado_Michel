import docx

doc_a = docx.Document(r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_antigravity.docx")
print("Antigravity P0137 palavras:", len(doc_a.paragraphs[137].text.split()))
print("Antigravity P0137 texto:\n", doc_a.paragraphs[137].text)

print("\n--- ABSTRACT ---")
print("Antigravity P0142 palavras:", len(doc_a.paragraphs[142].text.split()))
print("Antigravity P0142 texto:\n", doc_a.paragraphs[142].text)
