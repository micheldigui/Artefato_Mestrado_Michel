import re
import docx

doc = docx.Document(r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_merge.docx")

print("=== VERIFICACAO DE CITACAO DE FIGURAS NO TEXTO (1 a 28) ===")
# Mapear aparicoes de "Figura X" no texto corrido (excluindo a propria legenda "Figura X -")
citacoes_figuras = {i: [] for i in range(1, 29)}

for p_idx, p in enumerate(doc.paragraphs[293:]):
    actual_idx = 293 + p_idx
    txt = p.text.strip()
    if txt.startswith("Figura ") and ("-" in txt or ":" in txt):
        continue # E a propria legenda
    for n in range(1, 29):
        pattern = r'\bFigura\s+' + str(n) + r'\b'
        if re.search(pattern, txt):
            citacoes_figuras[n].append(actual_idx)

for n in range(1, 29):
    ocorrencias = citacoes_figuras[n]
    if ocorrencias:
        print(f"  Figura {n:02d}: Citada {len(ocorrencias)} vez(es) no texto (ex: P{ocorrencias[0]:04d})")
    else:
        print(f"  [ALERTA] Figura {n:02d}: NAO CITADA no texto corrido!")

print("\n=== VERIFICACAO DE CITACAO DE TABELAS NO TEXTO (1 a 8) ===")
citacoes_tabelas = {i: [] for i in range(1, 9)}
for p_idx, p in enumerate(doc.paragraphs[293:]):
    actual_idx = 293 + p_idx
    txt = p.text.strip()
    if txt.startswith("Tabela ") and ("-" in txt or ":" in txt):
        continue
    for n in range(1, 9):
        pattern = r'\bTabela\s+' + str(n) + r'\b'
        if re.search(pattern, txt):
            citacoes_tabelas[n].append(actual_idx)

for n in range(1, 9):
    ocorrencias = citacoes_tabelas[n]
    if ocorrencias:
        print(f"  Tabela {n:02d}: Citada {len(ocorrencias)} vez(es) no texto (ex: P{ocorrencias[0]:04d})")
    else:
        print(f"  [ALERTA] Tabela {n:02d}: NAO CITADA no texto corrido!")
