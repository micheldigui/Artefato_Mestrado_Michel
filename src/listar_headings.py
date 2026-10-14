import docx

doc_a = docx.Document(r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_antigravity.docx")
doc_c = docx.Document(r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_codex.docx")

print("--- HEADINGS CORPO ANTIGRAVITY (P >= 290) ---")
for i, p in enumerate(doc_a.paragraphs[290:]):
    if p.style.name.startswith("Heading"):
        print(f"P{290+i:04d} [{p.style.name}]: {p.text.strip()}")

print("\n--- HEADINGS CORPO CODEX (P >= 280) ---")
for i, p in enumerate(doc_c.paragraphs[280:]):
    if p.style.name.startswith("Heading"):
        print(f"P{280+i:04d} [{p.style.name}]: {p.text.strip()}")
