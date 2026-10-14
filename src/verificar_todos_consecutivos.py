import docx

doc = docx.Document(r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_merge.docx")

headings = []
for i, p in enumerate(doc.paragraphs[293:]):
    actual_i = 293 + i
    if p.style.name.startswith("Heading"):
        headings.append((actual_i, p.style.name, p.text.strip()))

print(f"Total de headings no corpo: {len(headings)}")

consecutivos = []
for k in range(len(headings) - 1):
    idx1, st1, txt1 = headings[k]
    idx2, st2, txt2 = headings[k+1]
    
    # Textos entre idx1 e idx2
    textos = []
    for j in range(idx1 + 1, idx2):
        t = doc.paragraphs[j].text.strip()
        if t:
            textos.append((j, t))
    if not textos:
        consecutivos.append((idx1, txt1, idx2, txt2))

print(f"Pares de headings consecutivos SEM texto intermediario no corpo: {len(consecutivos)}")
for idx1, txt1, idx2, txt2 in consecutivos:
    print(f"  P{idx1:04d} '{txt1}' -> P{idx2:04d} '{txt2}'")
