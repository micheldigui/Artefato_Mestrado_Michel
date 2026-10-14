import docx

doc = docx.Document(r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_merge.docx")
print("Palavras no corpo do RESUMO (P0137):", len(doc.paragraphs[137].text.split()))
print("Palavras no corpo do ABSTRACT (P0142):", len(doc.paragraphs[142].text.split()))
