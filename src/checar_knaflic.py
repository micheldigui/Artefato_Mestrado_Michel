import docx

doc_a = docx.Document(r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_antigravity.docx")
doc_m = docx.Document(r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_merge.docx")

print("--- BUSCANDO KNAFLIC EM ANTIGRAVITY ---")
for i, p in enumerate(doc_a.paragraphs[830:]):
    if "knaflic" in p.text.lower():
        print(f"Antigravity P{830+i:04d}: {p.text}")

print("\n--- BUSCANDO KNAFLIC EM MERGE ---")
for i, p in enumerate(doc_m.paragraphs[830:]):
    if "knaflic" in p.text.lower():
        print(f"Merge P{830+i:04d}: {p.text}")
