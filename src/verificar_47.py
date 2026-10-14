import docx

doc = docx.Document(r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_merge.docx")
print("Buscando 4.7 no documento...")
for i, p in enumerate(doc.paragraphs):
    if "4.7" in p.text:
        print(f"P{i:04d} [{p.style.name}]: {p.text[:100]}")
        for j in range(i, min(i+10, len(doc.paragraphs))):
            print(f"   P{j:04d}: {doc.paragraphs[j].text[:80]}")
