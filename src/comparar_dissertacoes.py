"""
Script para comparacao detalhada e estruturada entre:
- Dissertacao_Michel_qualificacao_antigravity.docx
- Dissertacao_Michel_qualificacao_codex.docx
"""

import docx
import difflib

PATH_AG = r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_antigravity.docx"
PATH_CD = r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_codex.docx"

doc_ag = docx.Document(PATH_AG)
doc_cd = docx.Document(PATH_CD)

print(f"Total paragrafos: AG={len(doc_ag.paragraphs)}, CD={len(doc_cd.paragraphs)}")
print(f"Total tabelas:    AG={len(doc_ag.tables)}, CD={len(doc_cd.tables)}")

headings_ag = [(i, p.text.strip()) for i, p in enumerate(doc_ag.paragraphs) if p.style.name.startswith('Heading')]
headings_cd = [(i, p.text.strip()) for i, p in enumerate(doc_cd.paragraphs) if p.style.name.startswith('Heading')]

print(f"\nHeadings encontrados: AG={len(headings_ag)}, CD={len(headings_cd)}")

min_len = min(len(headings_ag), len(headings_cd))
diff_h = 0
for k in range(min_len):
    i_ag, h_ag = headings_ag[k]
    i_cd, h_cd = headings_cd[k]
    if h_ag != h_cd:
        print(f"Diff H[{k}]: AG(P{i_ag}) '{h_ag[:40]}' != CD(P{i_cd}) '{h_cd[:40]}'")
        diff_h += 1

if diff_h == 0 and len(headings_ag) == len(headings_cd):
    print("Todos os headings sao identicos!")
elif len(headings_ag) != len(headings_cd):
    print(f"Quantidade de headings diferente: AG={len(headings_ag)}, CD={len(headings_cd)}")

# Comparar Resumo e Abstract
print("\n--- Comparacao de Resumo e Abstract ---")
print("AG P137 (Resumo):", doc_ag.paragraphs[137].text[:120], f"[{len(doc_ag.paragraphs[137].text.split())} palavras]")
print("CD P137 (Resumo):", doc_cd.paragraphs[137].text[:120], f"[{len(doc_cd.paragraphs[137].text.split())} palavras]")

print("\nAG P142 (Abstract):", doc_ag.paragraphs[142].text[:120], f"[{len(doc_ag.paragraphs[142].text.split())} palavras]")
print("CD P142 (Abstract):", doc_cd.paragraphs[142].text[:120], f"[{len(doc_cd.paragraphs[142].text.split())} palavras]")
