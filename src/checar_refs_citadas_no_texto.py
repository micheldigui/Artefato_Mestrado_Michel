import docx
import re

doc = docx.Document(r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_merge.docx")

p_refs_idx = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().upper() == "REFERÊNCIAS" or p.text.strip().upper() == "REFERENCIAS":
        p_refs_idx = i
        break

# Pega os paragrafos ate APENDICE
p_apendice_idx = len(doc.paragraphs)
for i in range(p_refs_idx, len(doc.paragraphs)):
    if "APÊNDICE" in doc.paragraphs[i].text.upper() or "APENDICE" in doc.paragraphs[i].text.upper():
        p_apendice_idx = i
        break

print(f"Referencias de P{p_refs_idx:04d} a P{p_apendice_idx:04d}")
refs = [p.text.strip() for p in doc.paragraphs[p_refs_idx+1:p_apendice_idx] if p.text.strip()]
print(f"Total de referencias listadas: {len(refs)}")

# Texto corrido
texto_corpo = "\n".join([p.text for p in doc.paragraphs[293:p_refs_idx]])

# Verificar quais autores principais de cada referencia aparecem no texto
nao_citados = []
for r in refs:
    # primeiro autor geralmente eh a primeira palavra em maiusculas
    m = re.match(r'^([A-ZÁÀÂÃÉÈÊÍÏÓÔÕÖÚÇÑ]+),', r)
    if m:
        sobrenome = m.group(1).title()
        # buscar sobrenome no texto
        if not re.search(r'\b' + re.escape(sobrenome) + r'\b', texto_corpo, re.IGNORECASE):
            nao_citados.append((sobrenome, r[:80]))

print(f"Referencias nao citadas diretamente pelo sobrenome: {len(nao_citados)}")
for s, r in nao_citados:
    print(f"  - {s}: {r}...")
