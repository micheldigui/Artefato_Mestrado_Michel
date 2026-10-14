"""
Inspecionar texto exato do Capitulo 6 no Codex
"""
import docx

doc_cd = docx.Document(r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_codex.docx")

p_c6 = None
for i, p in enumerate(doc_cd.paragraphs):
    if p.text.strip() == "6 CONSIDERAÇÕES PARCIAIS" or p.text.strip() == "6 CONSIDERAES PARCIAIS":
        p_c6 = i
        break

print(f"Capitulo 6 no Codex comeca em P{p_c6:04d}")
for i in range(p_c6, p_c6 + 10):
    if i < len(doc_cd.paragraphs):
        print(f"CD P{i:04d} [{doc_cd.paragraphs[i].style.name}]: {doc_cd.paragraphs[i].text}")
