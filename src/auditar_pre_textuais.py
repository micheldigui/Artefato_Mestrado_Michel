import docx

doc = docx.Document(r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_merge.docx")

print("=== AUDITORIA DOS ELEMENTOS PRE-TEXTUAIS ===")

# 1. Resumo e Abstract
for i in range(120, 160):
    p = doc.paragraphs[i]
    if p.text.strip() == "RESUMO":
        p_res = doc.paragraphs[i+1]
        print(f"Resumo encontrado em P{i+1:04d}: {len(p_res.text.split())} palavras")
        print(f"Primeiras 120 letras: {p_res.text[:120]}...")
    elif p.text.strip() == "ABSTRACT":
        p_abs = doc.paragraphs[i+1]
        print(f"Abstract encontrado em P{i+1:04d}: {len(p_abs.text.split())} palavras")
        print(f"Primeiras 120 letras: {p_abs.text[:120]}...")

# 2. Lista de Ilustracoes / Figuras
print("\n--- LISTA DE FIGURAS NO PRE-TEXTUAL ---")
figs_pre = []
for i in range(150, 200):
    p = doc.paragraphs[i]
    t = p.text.strip()
    if t.startswith("Figura "):
        figs_pre.append((i, t))
print(f"Total de figuras na Lista de Figuras pre-textual: {len(figs_pre)}")
for idx, f in figs_pre:
    print(f"  P{idx:04d}: {f}")

# 3. Lista de Quadros
print("\n--- LISTA DE QUADROS NO PRE-TEXTUAL ---")
quadros_pre = []
for i in range(170, 215):
    p = doc.paragraphs[i]
    t = p.text.strip()
    if t.startswith("Quadro "):
        quadros_pre.append((i, t))
print(f"Total de quadros na Lista de Quadros pre-textual: {len(quadros_pre)}")
for idx, q in quadros_pre[:5]:
    print(f"  P{idx:04d}: {q}")

# 4. Lista de Tabelas
print("\n--- LISTA DE TABELAS NO PRE-TEXTUAL ---")
tabs_pre = []
for i in range(180, 220):
    p = doc.paragraphs[i]
    t = p.text.strip()
    if t.startswith("Tabela "):
        tabs_pre.append((i, t))
print(f"Total de tabelas na Lista de Tabelas pre-textual: {len(tabs_pre)}")
for idx, t in tabs_pre:
    print(f"  P{idx:04d}: {t}")
