"""
Inspecionar diferencas de conteudo detalhadas no Capitulo 3 e Capitulo 5.
"""

import docx
import difflib

PATH_AG = r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_antigravity.docx"
PATH_CD = r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_codex.docx"

doc_ag = docx.Document(PATH_AG)
doc_cd = docx.Document(PATH_CD)

def get_chapter_paragraphs(doc, heading_start, heading_end):
    collecting = False
    paragraphs = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if p.style.name.startswith("Heading 1") and heading_start in t:
            collecting = True
            paragraphs.append((p.style.name, t))
            continue
        if collecting and p.style.name.startswith("Heading 1") and heading_end in t:
            break
        if collecting and t:
            paragraphs.append((p.style.name, t))
    return paragraphs

print("=== DIFERENCAS NO CAPITULO 3 (METODOLOGIA) ===")
p3_ag = get_chapter_paragraphs(doc_ag, "3 METODOLOGIA", "4 MODELO")
p3_cd = get_chapter_paragraphs(doc_cd, "3 METODOLOGIA", "4 MODELO")

matcher3 = difflib.SequenceMatcher(None, [t for _, t in p3_ag], [t for _, t in p3_cd])
for tag, i1, i2, j1, j2 in matcher3.get_opcodes():
    if tag != 'equal':
        print(f"\n[{tag.upper()}] AG({i1}:{i2}) vs CD({j1}:{j2}):")
        for idx in range(i1, min(i1+3, i2)):
            print(f"  - AG: {p3_ag[idx][1][:110]}...")
        for jdx in range(j1, min(j1+3, j2)):
            print(f"  + CD: {p3_cd[jdx][1][:110]}...")

print("\n\n=== ESTRUTURA DO CAPITULO 5 (RESULTADOS) ===")
p5_ag = get_chapter_paragraphs(doc_ag, "5 AN", "6 CONSIDERA")
p5_cd = get_chapter_paragraphs(doc_cd, "5 AN", "6 CONSIDERA")

headings_p5_ag = [t for s, t in p5_ag if s.startswith("Heading")]
headings_p5_cd = [t for s, t in p5_cd if s.startswith("Heading")]

print("\nHeadings Cap 5 no Antigravity:")
for h in headings_p5_ag:
    print("  AG:", h)

print("\nHeadings Cap 5 no Codex:")
for h in headings_p5_cd:
    print("  CD:", h)
