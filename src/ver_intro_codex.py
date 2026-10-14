"""
Inspecionar texto exato do Capitulo 1 no Codex
"""
import docx

doc_cd = docx.Document(r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_codex.docx")

p_intro = None
for i, p in enumerate(doc_cd.paragraphs):
    if p.text.strip() == "1 INTRODUÇÃO" or p.text.strip() == "1 INTRODUO":
        p_intro = i
        break

print(f"Introducao no Codex comeca em P{p_intro:04d}")
for i in range(p_intro, p_intro + 15):
    print(f"CD P{i:04d} [{doc_cd.paragraphs[i].style.name}]: {doc_cd.paragraphs[i].text}")
