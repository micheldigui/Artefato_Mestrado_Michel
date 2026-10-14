import docx

doc = docx.Document(r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_merge.docx")
print("P0134:", doc.paragraphs[134].text[:60])
print("P0135:", doc.paragraphs[135].text[:60])
print("P0136:", doc.paragraphs[136].text[:80])
print("Palavras em P0136:", len(doc.paragraphs[136].text.split()))
print("P0137:", doc.paragraphs[137].text[:60])
print("---")
print("P0139:", doc.paragraphs[139].text[:60])
print("P0140:", doc.paragraphs[140].text[:60])
print("P0141:", doc.paragraphs[141].text[:80])
print("Palavras em P0141:", len(doc.paragraphs[141].text.split()))
print("P0142:", doc.paragraphs[142].text[:60])
