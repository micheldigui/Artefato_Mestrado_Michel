import re
import docx

doc = docx.Document(r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_merge.docx")

# Localizar REFERENCIAS
p_refs_idx = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().upper() == "REFERÊNCIAS" or p.text.strip().upper() == "REFERENCIAS":
        p_refs_idx = i
        break

print(f"Secao REFERENCIAS inicia em P{p_refs_idx:04d}")

# Juntar texto de todas as referencias
texto_refs = "\n".join([p.text.strip() for p in doc.paragraphs[p_refs_idx:] if p.text.strip()])

# Extrair autores citados no corpo (P0293 ate p_refs_idx)
autores_citados = set()
padroes = [
    r'\(([A-Z][a-záàâãéèêíïóôõöúçñ]+(?:\s+et\s+al\.)?),?\s+(\d{4})\)',
    r'\b([A-Z][a-záàâãéèêíïóôõöúçñ]+)\s+et\s+al\.\s+\((\d{4})\)',
    r'\b([A-Z][a-záàâãéèêíïóôõöúçñ]+(?:\s+e\s+[A-Z][a-záàâãéèêíïóôõöúçñ]+)?)\s+\((\d{4})\)',
]

for p in doc.paragraphs[293:p_refs_idx]:
    txt = p.text
    for pat in padroes:
        for match in re.finditer(pat, txt):
            autor = match.group(1).replace(" et al.", "").strip()
            # Se tiver ' e '
            if " e " in autor:
                for a in autor.split(" e "):
                    autores_citados.add(a.strip())
            else:
                autores_citados.add(autor)

print(f"\nTotal de autores/chaves de citacao unicos identificados no texto: {len(autores_citados)}")

# Checar quais estao presentes nas REFERENCIAS
ausentes = []
presentes = []
for autor in sorted(autores_citados):
    # ignorar palavras comuns que possam casar por acaso
    if autor.lower() in ["figura", "tabela", "quadro", "secao", "capitulo", "apendice", "lei", "decreto"]:
        continue
    if re.search(r'\b' + re.escape(autor) + r'\b', texto_refs, re.IGNORECASE):
        presentes.append(autor)
    else:
        ausentes.append(autor)

print(f"Autores presentes nas REFERENCIAS: {len(presentes)}")
print(f"Autores ausentes nas REFERENCIAS: {len(ausentes)}")
if ausentes:
    print("\n[ALERTA] Os seguintes autores citados no texto nao foram localizados na lista de Referencias:")
    for a in ausentes:
        print(f"  - {a}")
