import os
import re
import docx

PATH_MERGE = r"f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_merge.docx"

doc = docx.Document(PATH_MERGE)

# Localizar o inicio do corpo (P com 1 INTRODUCAO)
p_intro_idx = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().upper() == "1 INTRODUÇÃO" or p.text.strip().upper() == "1 INTRODUO":
        p_intro_idx = i
        break

print(f"Inicio do corpo textual: P{p_intro_idx:04d} ('{doc.paragraphs[p_intro_idx].text.strip()}')")

# Obter todos os headings do corpo
headings_corpo = []
for i in range(p_intro_idx, len(doc.paragraphs)):
    p = doc.paragraphs[i]
    if p.style.name.startswith("Heading"):
        headings_corpo.append((i, p.style.name, p.text.strip()))

print(f"\nTotal de headings encontrados no corpo: {len(headings_corpo)}")

# Verificar headings consecutivos sem texto no corpo
falhas = 0
for idx in range(len(headings_corpo) - 1):
    curr_i, curr_st, curr_txt = headings_corpo[idx]
    next_i, next_st, next_txt = headings_corpo[idx + 1]
    
    # Textos entre os dois headings
    textos_meio = []
    for j in range(curr_i + 1, next_i):
        t = doc.paragraphs[j].text.strip()
        if t:
            textos_meio.append((j, t))
            
    if not textos_meio:
        print(f"[ALERTA FALHA] P{curr_i:04d} [{curr_st}] '{curr_txt}' seguido direto por P{next_i:04d} [{next_st}] '{next_txt}' SEM TEXTO!")
        falhas += 1
    else:
        # Se ha apenas legenda ou tabela
        txt_apenas = [t[1] for t in textos_meio if not t[1].startswith("Figura ") and not t[1].startswith("Tabela ") and not t[1].startswith("Quadro ") and not t[1].startswith("Fonte:")]
        if not txt_apenas:
            print(f"[AVISO] P{curr_i:04d} '{curr_txt}' tem apenas ilustracao antes de P{next_i:04d} '{next_txt}'!")

if falhas == 0:
    print("\n[SUCESSO ABSOLUTO] NENHUM heading no corpo da dissertacao esta sem texto intermediario!")
else:
    print(f"\n[FALHAS ENCONTRADAS] Total de headings sem texto intermediario: {falhas}")
