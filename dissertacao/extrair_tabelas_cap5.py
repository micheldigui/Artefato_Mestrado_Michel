import docx

doc_path = r'f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_resultados_parciais_20261013_antigravity.docx'
doc = docx.Document(doc_path)

with open(r'f:\PROJETOS\ARTEFATO\dissertacao\tabelas_capitulo5.txt', 'w', encoding='utf-8') as f:
    for t_idx in range(16, len(doc.tables)):
        t = doc.tables[t_idx]
        f.write(f"\n{'='*70}\nTABELA DOCX ÍNDICE {t_idx} ({len(t.rows)} linhas, {len(t.columns)} colunas)\n{'='*70}\n")
        for r_idx, r in enumerate(t.rows):
            vals = [c.text.strip().replace('\n', ' ') for c in r.cells]
            f.write(f"L{r_idx}: {vals}\n")

print("Tabelas do capítulo 5 extraídas com sucesso!")
