import docx
import re
import json

doc_path = r'f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_resultados_parciais_20261013_entrega.docx'
doc = docx.Document(doc_path)

# Encontrar relacionamento de imagens
part = doc.part
rels = part.rels

figure_data = []

# Varrer parágrafos
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    m = re.match(r'^Figura\s+(\d+)\s*[-–—:]\s*(.*)', txt, re.IGNORECASE)
    if m:
        fig_num = int(m.group(1))
        fig_title = m.group(2).strip()
        
        # Procurar parágrafo com imagem logo abaixo (geralmente i+1 ou i+2)
        img_info = None
        img_p_idx = None
        for offset in range(1, 4):
            if i + offset < len(doc.paragraphs):
                p_cand = doc.paragraphs[i + offset]
                xml = p_cand._p.xml
                if 'blip' in xml or 'graphic' in xml:
                    # extrair r:embed
                    embed_m = re.findall(r'r:embed="([^"]+)"', xml)
                    if embed_m:
                        r_id = embed_m[0]
                        target_ref = rels[r_id].target_ref
                        img_info = target_ref
                        img_p_idx = i + offset
                        break
        
        # Procurar fonte (geralmente abaixo da imagem)
        fonte_text = ""
        for offset in range(1, 5):
            if i + offset < len(doc.paragraphs):
                p_cand = doc.paragraphs[i + offset]
                if p_cand.text.strip().lower().startswith('fonte'):
                    fonte_text = p_cand.text.strip()
                    break

        # Parágrafo anterior (introdução) e posterior (análise)
        p_ant = doc.paragraphs[i-1].text.strip() if i > 0 else ""
        p_pos = ""
        # achar primeiro parágrafo de texto corrido após a fonte
        for offset in range(1, 6):
            if i + offset < len(doc.paragraphs):
                p_cand = doc.paragraphs[i + offset]
                ptxt = p_cand.text.strip()
                if ptxt and not ptxt.lower().startswith('fonte') and not any(r in p_cand._p.xml for r in ['graphic', 'blip']):
                    p_pos = ptxt
                    break

        figure_data.append({
            'num': fig_num,
            'caption_p': i,
            'title': fig_title,
            'img_p': img_p_idx,
            'img_target': img_info,
            'fonte': fonte_text,
            'context_ant': p_ant[:120],
            'context_pos': p_pos[:150]
        })

print(f"Total de figuras mapeadas no texto: {len(figure_data)}")
with open('mapeamento_figuras_entrega.json', 'w', encoding='utf-8') as f:
    json.dump(figure_data, f, indent=2, ensure_ascii=False)

print("Mapeamento salvo em mapeamento_figuras_entrega.json")
