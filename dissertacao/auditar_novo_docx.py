import docx
import os
import re
import json

doc_path = r'f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_resultados_parciais_20261013_entrega.docx'
doc = docx.Document(doc_path)

print(f"Total de parágrafos: {len(doc.paragraphs)}")
print(f"Total de tabelas: {len(doc.tables)}")

# 1. Localizar todas as Figuras, Quadros e Tabelas
captions = []
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    m = re.match(r'^(Figura|Quadro|Tabela)\s+(\d+)\s*[-–—:]\s*(.*)', txt, re.IGNORECASE)
    if m:
        captions.append({
            'p_idx': i,
            'type': m.group(1).capitalize(),
            'num': int(m.group(2)),
            'title': m.group(3).strip()
        })

print(f"Total de legendas encontradas: {len(captions)}")

with open('auditoria_novo_docx_captions.json', 'w', encoding='utf-8') as f:
    json.dump(captions, f, indent=2, ensure_ascii=False)

print("Legendas salvas em auditoria_novo_docx_captions.json")
