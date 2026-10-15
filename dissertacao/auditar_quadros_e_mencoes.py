import docx
import re
import json

doc_path = r'f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_resultados_parciais_20261013_entrega.docx'
doc = docx.Document(doc_path)

print("=== 1. CONFERÊNCIA EXAUSTIVA DOS 16 QUADROS ===")
quadros_detalhes = []
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    m = re.match(r'^Quadro\s+(\d+)\s*[-–—:]\s*(.*)', txt, re.IGNORECASE)
    if m and i > 200:
        q_num = int(m.group(1))
        q_title = m.group(2).strip()
        p_ant = doc.paragraphs[i-1].text.strip() if i > 0 else ""
        
        # Achar fonte
        fonte = ""
        for off in range(1, 4):
            if i + off < len(doc.paragraphs):
                cand = doc.paragraphs[i + off]
                if cand.text.strip().lower().startswith('fonte'):
                    fonte = cand.text.strip()
                    break
        
        # Achar texto posterior
        p_pos = ""
        for off in range(1, 6):
            if i + off < len(doc.paragraphs):
                cand = doc.paragraphs[i + off]
                ctxt = cand.text.strip()
                if ctxt and not ctxt.lower().startswith('fonte') and not any(r in cand._p.xml for r in ['graphic', 'blip']):
                    p_pos = ctxt
                    break
                    
        # Checar se o quadro é mencionado no texto
        mencoes = []
        for p_idx in range(200, len(doc.paragraphs)):
            if p_idx == i:
                continue
            if re.search(r'\bQuadro\s+' + str(q_num) + r'\b', doc.paragraphs[p_idx].text, re.IGNORECASE):
                mencoes.append(p_idx)

        quadros_detalhes.append({
            'num': q_num,
            'p_idx': i,
            'title': q_title,
            'fonte': fonte,
            'p_ant': p_ant[:100],
            'p_pos': p_pos[:100],
            'mencoes': mencoes
        })

print(f"Total de quadros no corpo: {len(quadros_detalhes)}")
for q in quadros_detalhes:
    print(f"Quadro {q['num']:02d} (P{q['p_idx']}): '{q['title'][:40]}' | Mencionada em P: {q['mencoes']}")

# 2. CONFERÊNCIA DE MENÇÕES DE FIGURAS NO TEXTO (BUSCA POR INCONSISTÊNCIAS)
print("\n=== 2. VERIFICAÇÃO DE MENÇÕES A FIGURAS NO CORPO DO TEXTO ===")
# Procurar todas as menções 'Figura X' em parágrafos do corpo do texto
fig_mismatches = []
for p_idx in range(700, 850):
    p_text = doc.paragraphs[p_idx].text
    for m in re.finditer(r'\bFigura\s+(\d+)\b', p_text):
        f_num = int(m.group(1))
        # Salvar para checar o contexto
        fig_mismatches.append((p_idx, f_num, p_text[:110]))

for p_i, fn, ctx in fig_mismatches:
    print(f"[P{p_i}] Cita Figura {fn}: {ctx}...")
