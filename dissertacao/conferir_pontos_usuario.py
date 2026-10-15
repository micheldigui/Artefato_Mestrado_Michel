import docx

doc = docx.Document('Dissertacao_Michel_qualificacao_resultados_parciais_20261013_antigravity.docx')

print("="*80)
print("1. TABELA 3 E CONTEUDO")
print("="*80)
print(f"Legenda P696: {doc.paragraphs[696].text}")
print(f"Fonte P697: {doc.paragraphs[697].text}")
# Procurar tabela correspondente
for t_idx, t in enumerate(doc.tables):
    txt_all = " ".join(c.text for row in t.rows for c in row.cells)
    if "2023" in txt_all and "2024" in txt_all and "2025" in txt_all and "12.410" in txt_all or "6.578" in txt_all or "Reserva" in txt_all:
        print(f"Tabela encontrada: indice {t_idx} com {len(t.rows)} linhas e {len(t.columns)} colunas")
        for r_i, r in enumerate(t.rows):
            row_vals = [c.text.strip().replace('\n', ' ') for c in r.cells]
            print(f"   L{r_i}: {row_vals}")
        break

print("\n" + "="*80)
print("2. FIGURA 16 E PARAGRAFOS (P700 a P704)")
print("="*80)
for p in range(700, 705):
    print(f"P{p}: {doc.paragraphs[p].text}")

print("\n" + "="*80)
print("3. FIGURA 17 E PARAGRAFOS (P705 a P710)")
print("="*80)
for p in range(705, 711):
    print(f"P{p}: {doc.paragraphs[p].text}")

print("\n" + "="*80)
print("4. FIGURA 18 E PARAGRAFOS (P720 a P728)")
print("="*80)
for p in range(720, 729):
    print(f"P{p}: {doc.paragraphs[p].text}")

print("\n" + "="*80)
print("5. FIGURA 23 E PARAGRAFOS (P782 a P788)")
print("="*80)
for p in range(782, 789):
    print(f"P{p}: {doc.paragraphs[p].text}")
