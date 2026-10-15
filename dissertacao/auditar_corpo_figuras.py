import json
import os
import docx
import re
from PIL import Image

doc_path = r'f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_resultados_parciais_20261013_entrega.docx'
doc = docx.Document(doc_path)
media_dir = r'f:\PROJETOS\ARTEFATO\scratch_entrega_imgs\word\media'

with open('mapeamento_figuras_entrega.json', 'r', encoding='utf-8') as f:
    figs = json.load(f)

# Filtrar apenas as do corpo do texto (onde caption_p > 200)
body_figs = [f for f in figs if f['caption_p'] > 200]
print(f"Total de figuras no corpo do texto: {len(body_figs)}")

with open('auditoria_detalhada_corpo_texto.txt', 'w', encoding='utf-8') as out:
    out.write("=== AUDITORIA MINUCIOSA DAS FIGURAS NO CORPO DO TEXTO ===\n\n")
    for fig in body_figs:
        n = fig['num']
        p_idx = fig['caption_p']
        title = fig['title']
        img_target = fig['img_target']
        fonte = fig['fonte']
        
        # Obter detalhes da imagem
        img_filename = os.path.basename(img_target) if img_target else "SEM_IMAGEM"
        img_path = os.path.join(media_dir, img_filename) if img_target else ""
        img_dims = "N/A"
        img_size = 0
        if os.path.exists(img_path):
            img_size = os.path.getsize(img_path)
            try:
                im = Image.open(img_path)
                img_dims = f"{im.size[0]}x{im.size[1]} ({im.format})"
            except:
                pass
                
        # Obter texto completo do parágrafo anterior e posterior
        p_ant_txt = doc.paragraphs[p_idx - 1].text.strip() if p_idx > 0 else ""
        # achar p_pos real
        p_pos_txt = ""
        for offset in range(1, 6):
            if p_idx + offset < len(doc.paragraphs):
                cand = doc.paragraphs[p_idx + offset]
                ctxt = cand.text.strip()
                if ctxt and not ctxt.lower().startswith('fonte') and not any(r in cand._p.xml for r in ['graphic', 'blip']):
                    p_pos_txt = ctxt
                    break
                    
        out.write(f"--------------------------------------------------------------------------------\n")
        out.write(f"FIGURA {n} (Parágrafo {p_idx})\n")
        out.write(f"Legenda: Figura {n} - {title}\n")
        out.write(f"Arquivo Mídia Docx: {img_filename} | Dimensões: {img_dims} | Tamanho: {img_size} bytes\n")
        out.write(f"Linha de Fonte: {fonte}\n")
        out.write(f"Contexto Anterior (P{p_idx-1}):\n{p_ant_txt}\n\n")
        out.write(f"Contexto Posterior:\n{p_pos_txt}\n\n")

print("Relatório salvo em auditoria_detalhada_corpo_texto.txt")
