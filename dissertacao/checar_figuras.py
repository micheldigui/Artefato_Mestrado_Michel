import docx
import re

doc = docx.Document('Dissertacao_Michel_qualificacao_resultados_parciais_20261013_antigravity.docx')

print("=== CONFERÊNCIA DETALHADA DAS 28 FIGURAS ===")
for n in range(1, 29):
    leg_p = None
    leg_title = ''
    for i in range(400, 866):
        txt = doc.paragraphs[i].text.strip()
        m = re.match(r'^Figura\s+' + str(n) + r'\s*[-–—:]\s*(.*)', txt, re.IGNORECASE)
        if m:
            leg_p = i
            leg_title = m.group(1).strip()
            break
    
    mencoes = []
    for i in range(200, 866):
        if i == leg_p:
            continue
        txt = doc.paragraphs[i].text
        if re.search(r'\bFigura\s+' + str(n) + r'\b', txt, re.IGNORECASE):
            mencoes.append(i)
            
    print(f"Figura {n:02d} (Legenda em P{leg_p}: '{leg_title}') -> Mencionada em P: {mencoes}")
