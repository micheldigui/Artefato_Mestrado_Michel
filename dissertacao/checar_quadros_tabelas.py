import docx
import re

doc = docx.Document('Dissertacao_Michel_qualificacao_resultados_parciais_20261013_antigravity.docx')

print("=== CONFERÊNCIA DETALHADA DOS QUADROS (1 A 16) ===")
for n in range(1, 17):
    leg_p = None
    leg_title = ''
    for i in range(250, 866):
        txt = doc.paragraphs[i].text.strip()
        m = re.match(r'^Quadro\s+' + str(n) + r'\s*[-–—:]\s*(.*)', txt, re.IGNORECASE)
        if m:
            leg_p = i
            leg_title = m.group(1).strip()
            break
    
    mencoes = []
    for i in range(200, 866):
        if i == leg_p:
            continue
        txt = doc.paragraphs[i].text
        if re.search(r'\bQuadro\s+' + str(n) + r'\b', txt, re.IGNORECASE):
            mencoes.append(i)
            
    print(f"Quadro {n:02d} (Legenda em P{leg_p}: '{leg_title[:40]}') -> Mencionada em P: {mencoes}")

print("\n=== CONFERÊNCIA DETALHADA DAS TABELAS (1 A 8) ===")
for n in range(1, 9):
    leg_p = None
    leg_title = ''
    for i in range(250, 866):
        txt = doc.paragraphs[i].text.strip()
        m = re.match(r'^Tabela\s+' + str(n) + r'\s*[-–—:]\s*(.*)', txt, re.IGNORECASE)
        if m:
            leg_p = i
            leg_title = m.group(1).strip()
            break
    
    mencoes = []
    for i in range(200, 866):
        if i == leg_p:
            continue
        txt = doc.paragraphs[i].text
        if re.search(r'\bTabela\s+' + str(n) + r'\b', txt, re.IGNORECASE):
            mencoes.append(i)
            
    print(f"Tabela {n:02d} (Legenda em P{leg_p}: '{leg_title[:40]}') -> Mencionada em P: {mencoes}")
