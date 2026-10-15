import docx

doc = docx.Document('Dissertacao_Michel_qualificacao_resultados_parciais_20261013_antigravity.docx')

with open('conferir_tabela3_e_figuras.txt', 'w', encoding='utf-8') as f:
    f.write("=== PARÁGRAFOS 690 A 730 ===\n")
    for i in range(690, 730):
        f.write(f"P{i}: {doc.paragraphs[i].text}\n")
    
    f.write("\n=== TODAS AS TABELAS DO DOCX ===\n")
    for t_idx, t in enumerate(doc.tables):
        f.write(f"\n--- TABELA {t_idx} ({len(t.rows)} linhas, {len(t.columns)} colunas) ---\n")
        for r_idx, r in enumerate(t.rows[:6]):
            vals = [c.text.strip().replace('\n', ' ') for c in r.cells]
            f.write(f"  L{r_idx}: {vals}\n")

print("Arquivo gerado com sucesso!")
