import json
import os
import docx
import re
from PIL import Image

doc_path = r'f:\PROJETOS\ARTEFATO\dissertacao\Dissertacao_Michel_qualificacao_resultados_parciais_20261013_entrega.docx'
doc = docx.Document(doc_path)
media_dir = r'f:\PROJETOS\ARTEFATO\scratch_entrega_imgs\word\media'

part = doc.part
rels = part.rels

report = {
    'figuras': [],
    'quadros': [],
    'tabelas': [],
    'divergencias': [],
    'avisos': []
}

# 1. MAPEAR FIGURAS
fig_count = 0
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    m = re.match(r'^Figura\s+(\d+)\s*[-–—:]\s*(.*)', txt, re.IGNORECASE)
    if m and i > 200: # corpo do texto
        fig_num = int(m.group(1))
        fig_title = m.group(2).strip()
        
        # Achar imagem
        img_info = None
        img_p = None
        for off in range(1, 4):
            if i + off < len(doc.paragraphs):
                cand = doc.paragraphs[i + off]
                xml = cand._p.xml
                if 'blip' in xml or 'graphic' in xml:
                    embed_m = re.findall(r'r:embed="([^"]+)"', xml)
                    if embed_m:
                        img_info = rels[embed_m[0]].target_ref
                        img_p = i + off
                        break
        
        # Achar fonte
        fonte = ""
        for off in range(1, 5):
            if i + off < len(doc.paragraphs):
                cand = doc.paragraphs[i + off]
                if cand.text.strip().lower().startswith('fonte'):
                    fonte = cand.text.strip()
                    break
                    
        p_ant = doc.paragraphs[i-1].text.strip() if i > 0 else ""
        p_pos = ""
        for off in range(1, 6):
            if i + off < len(doc.paragraphs):
                cand = doc.paragraphs[i + off]
                ctxt = cand.text.strip()
                if ctxt and not ctxt.lower().startswith('fonte') and not any(r in cand._p.xml for r in ['graphic', 'blip']):
                    p_pos = ctxt
                    break
                    
        report['figuras'].append({
            'num': fig_num,
            'p_idx': i,
            'title': fig_title,
            'img_file': os.path.basename(img_info) if img_info else "SEM_IMAGEM",
            'fonte': fonte,
            'p_ant': p_ant,
            'p_pos': p_pos
        })

# 2. MAPEAR QUADROS
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    m = re.match(r'^Quadro\s+(\d+)\s*[-–—:]\s*(.*)', txt, re.IGNORECASE)
    if m and i > 200:
        q_num = int(m.group(1))
        q_title = m.group(2).strip()
        
        fonte = ""
        for off in range(1, 5):
            if i + off < len(doc.paragraphs):
                cand = doc.paragraphs[i + off]
                if cand.text.strip().lower().startswith('fonte'):
                    fonte = cand.text.strip()
                    break
                    
        p_ant = doc.paragraphs[i-1].text.strip() if i > 0 else ""
        p_pos = ""
        for off in range(1, 6):
            if i + off < len(doc.paragraphs):
                cand = doc.paragraphs[i + off]
                ctxt = cand.text.strip()
                if ctxt and not ctxt.lower().startswith('fonte') and not any(r in cand._p.xml for r in ['graphic', 'blip']):
                    p_pos = ctxt
                    break
                    
        report['quadros'].append({
            'num': q_num,
            'p_idx': i,
            'title': q_title,
            'fonte': fonte,
            'p_ant': p_ant,
            'p_pos': p_pos
        })

# 3. MAPEAR TABELAS
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    m = re.match(r'^Tabela\s+(\d+)\s*[-–—:]\s*(.*)', txt, re.IGNORECASE)
    if m and i > 200:
        t_num = int(m.group(1))
        t_title = m.group(2).strip()
        
        fonte = ""
        for off in range(1, 5):
            if i + off < len(doc.paragraphs):
                cand = doc.paragraphs[i + off]
                if cand.text.strip().lower().startswith('fonte'):
                    fonte = cand.text.strip()
                    break
                    
        p_ant = doc.paragraphs[i-1].text.strip() if i > 0 else ""
        p_pos = ""
        for off in range(1, 6):
            if i + off < len(doc.paragraphs):
                cand = doc.paragraphs[i + off]
                ctxt = cand.text.strip()
                if ctxt and not ctxt.lower().startswith('fonte') and not any(r in cand._p.xml for r in ['graphic', 'blip']):
                    p_pos = ctxt
                    break
                    
        report['tabelas'].append({
            'num': t_num,
            'p_idx': i,
            'title': t_title,
            'fonte': fonte,
            'p_ant': p_ant,
            'p_pos': p_pos
        })

with open('auditoria_elementos_completos.json', 'w', encoding='utf-8') as f:
    json.dump(report, f, indent=2, ensure_ascii=False)

print(f"Figuras mapeadas: {len(report['figuras'])}")
print(f"Quadros mapeados: {len(report['quadros'])}")
print(f"Tabelas mapeadas: {len(report['tabelas'])}")
