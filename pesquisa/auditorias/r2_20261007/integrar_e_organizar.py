"""Integra as células produzidas por Artifact Tool preservando o pacote XLSX original."""
from pathlib import Path
from copy import deepcopy
import zipfile, json, hashlib, shutil
from collections import Counter
from lxml import etree as E
import openpyxl

ROOT=Path(r'F:\PRODUTO_MESTRADO'); OUT=ROOT/'pesquisa/auditorias/r2_20261007'
BOOK=ROOT/'pesquisa/matriz_artigos.xlsx'
plan=json.loads((OUT/'plano_organizacao.json').read_text(encoding='utf-8'))
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'; ns={'s':NS}
tag=lambda x:f'{{{NS}}}{x}'
with zipfile.ZipFile(OUT/'matriz_artigos_antes_da_inclusao.xlsx') as z:
    original={n:z.read(n) for n in z.namelist()}
with zipfile.ZipFile(OUT/'matriz_edicao_artifact.xlsx') as z:
    authored={n:z.read(n) for n in z.namelist()}
strings=[]
if 'xl/sharedStrings.xml' in authored:
    for si in E.fromstring(authored['xl/sharedStrings.xml']):
        strings.append(''.join(si.xpath('.//s:t/text()',namespaces=ns)))
styles=E.fromstring(original['xl/styles.xml']); xfs=styles.find(tag('cellXfs'))
style_cache={}
def wrapped_style(s):
    if s in style_cache:return style_cache[s]
    xf=deepcopy(xfs[int(s)]); alignment=xf.find(tag('alignment'))
    if alignment is None: alignment=E.SubElement(xf,tag('alignment'))
    alignment.set('wrapText','1'); xf.set('applyAlignment','1')
    found=next((i for i,x in enumerate(xfs) if E.tostring(x)==E.tostring(xf)),None)
    if found is None:found=len(xfs);xfs.append(xf)
    style_cache[s]=str(found);return str(found)
def converted(c,style):
    c=deepcopy(c)
    if c.get('t')=='s':
        text=strings[int(c.find(tag('v')).text)]
        for child in list(c):c.remove(child)
        c.set('t','inlineStr');is_=E.SubElement(c,tag('is'));t=E.SubElement(is_,tag('t'));t.text=text
        t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
    c.set('s',style);return c
def serialize(r):return E.tostring(r,xml_declaration=True,encoding='UTF-8',standalone=True)
main=E.fromstring(original['xl/worksheets/sheet1.xml']); a_main=E.fromstring(authored['xl/worksheets/sheet1.xml'])
summary=E.fromstring(original['xl/worksheets/sheet2.xml']); a_summary=E.fromstring(authored['xl/worksheets/sheet2.xml'])
data=main.find(tag('sheetData')); template=data.find(f"{tag('row')}[@r='77']")
template_styles={''.join(filter(str.isalpha,c.get('r'))):c.get('s','0') for c in template}
for r in range(78,83):
    source=a_main.find(f"{tag('sheetData')}/{tag('row')}[@r='{r}']")
    row=E.Element(tag('row'),attrib=dict(source.attrib))
    for cell in source:
        col=''.join(filter(str.isalpha,cell.get('r')))
        row.append(converted(cell,wrapped_style(template_styles.get(col,'0'))))
    data.append(row)
if main.find(tag('dimension')) is not None:
    main.find(tag('dimension')).set('ref','A1:R82')
# A tabela estava em A1:R76, excluindo A077. Inclui agora todas as 81 referências.
for name,b in list(original.items()):
    if name.startswith('xl/tables/') and name.endswith('.xml'):
        t=E.fromstring(b)
        if t.get('name')=='artigos':
            t.set('ref','A1:R82')
            if t.find(tag('autoFilter')) is not None:t.find(tag('autoFilter')).set('ref','A1:R82')
            original[name]=serialize(t)
if main.find(tag('autoFilter')) is not None:main.find(tag('autoFilter')).set('ref','A1:R82')
sdata=summary.find(tag('sheetData')); adata=a_summary.find(tag('sheetData'))
for addr in ['B1','B2','B3','B4','B6','A6']:
    source=adata.find(f".//{tag('c')}[@r='{addr}']")
    old=sdata.find(f".//{tag('c')}[@r='{addr}']")
    target=sdata.find(f"{tag('row')}[@r='{''.join(filter(str.isdigit,addr))}']")
    style=old.get('s','0') if old is not None else '0'
    if old is not None:target.remove(old)
    new=converted(source,wrapped_style(style) if addr=='A6' else style)
    target.append(new)
    ordered=sorted(list(target),key=lambda c:openpyxl.utils.column_index_from_string(''.join(filter(str.isalpha,c.get('r')))))
    for c in list(target):target.remove(c)
    target.extend(ordered)
    if addr=='A6':target.set('ht','30');target.set('customHeight','1')
for r in range(82,89):
    source=adata.find(f"{tag('row')}[@r='{r}']")
    if source is None:continue
    row=E.Element(tag('row'),attrib=dict(source.attrib))
    templ=sdata.find(f"{tag('row')}[@r='{'8' if r==83 else '9'}']")
    stylemap={''.join(filter(str.isalpha,c.get('r'))):c.get('s','0') for c in templ}
    for c in source:
        col=''.join(filter(str.isalpha,c.get('r')))
        row.append(converted(c,wrapped_style(stylemap.get(col,'0'))))
    if r==82:row.set('ht','30');row.set('customHeight','1')
    # Linhas 82–88 eram vazias, mas podem conter formatação: preservar os demais atributos/células.
    existing=sdata.find(f"{tag('row')}[@r='{r}']")
    if existing is not None:
        for c in existing:
            if not any(x.get('r')==c.get('r') for x in row):row.append(deepcopy(c))
        sdata.remove(existing)
    sdata.append(row)
ordered=sorted(list(sdata),key=lambda r:int(r.get('r')))
for r in list(sdata):sdata.remove(r)
sdata.extend(ordered)
xfs.set('count',str(len(xfs)))
original['xl/styles.xml']=serialize(styles)
original['xl/worksheets/sheet1.xml']=serialize(main)
original['xl/worksheets/sheet2.xml']=serialize(summary)
tmp=OUT/'matriz_integrada.xlsx'
with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as z:
    for n,b in original.items():z.writestr(n,b)
before=openpyxl.load_workbook(OUT/'matriz_artigos_antes_da_inclusao.xlsx')
after=openpyxl.load_workbook(tmp)
allowed={'Página1':{f'{openpyxl.utils.get_column_letter(c)}{r}' for r in range(78,83) for c in range(1,19)},
         'Página2':{'A6','B1','B2','B3','B4','B6'}|{f'{openpyxl.utils.get_column_letter(c)}{r}' for r in range(82,89) for c in range(1,6)}}
for s in before:
    for row in s:
        for c in row:
            if c.coordinate in allowed[s.title]:continue
            new=after[s.title][c.coordinate]
            assert c.value==new.value,(s.title,c.coordinate)
            assert c._style==new._style,(s.title,c.coordinate,'style')
refs=[r for r in list(after['Página1'].values)[1:] if r[0]]
assert len(refs)==81 and len({r[0] for r in refs})==81
assert after['Página1'].tables['artigos'].ref=='A1:R82'
cached=openpyxl.load_workbook(tmp,data_only=True)
assert [cached['Página2'][f'B{i}'].value for i in [1,2,3,4,6]]==[24,23,32,2,81]
assert hashlib.sha256((ROOT/'dissertacao/Dissertacao_Michel_qualificacao_resultados_parciais_20261006.docx').read_bytes()).hexdigest()==plan['sha256_doc']
shutil.copy2(tmp,BOOK)
for p in plan['papers']:
    src=ROOT/'pesquisa'/p['source']; dest=Path(p['destination'])
    resolved=dest.resolve(); assert resolved.is_relative_to((ROOT/'pesquisa').resolve())
    dest.parent.mkdir(parents=True,exist_ok=True)
    if dest.exists():assert hashlib.sha256(dest.read_bytes()).hexdigest()==p['sha256']
    else:shutil.move(str(src),str(dest))
    assert hashlib.sha256(dest.read_bytes()).hexdigest()==p['sha256']
    ref=f"{p['authors']}. {p['title']}. {p['journal']}, {p['year']}. DOI: https://doi.org/{p['doi']}."
    p['reference']=ref
    summary_path=Path(p['summary']);summary_path.parent.mkdir(parents=True,exist_ok=True)
    text=f"# {p['id']} — {p['title']}\n\n## Referência\n\n{ref}\n\n## Classificação\n\n{p['category']}. Uso principal: seção 5.7; apoio à seção 2.4.1 e ao método em 3.5.\n\n## Evidência conferida no PDF\n\n{p['evidence']}\n\nLocalização: {p['pages']}\n\n## Ideia a aproveitar\n\n{p['idea']}\n\n## Texto possível\n\n{p['sentence']}\n\n## Limites da utilização\n\n{p['limits']}\n\n## Arquivo\n\n[{dest.name}]({dest.as_posix()})\n\nSHA-256: {p['sha256']}\n"
    summary_path.write_text(text,encoding='utf-8')
    dest.with_suffix('.md').write_text(f"# {p['title']}\n\n{ref}\n\nClassificação: {p['category']}.\n\n[Ficha de leitura e uso na dissertação]({summary_path.as_posix()})\n",encoding='utf-8')
counts=Counter(str(r[17]).replace('\u00a0',' ') for r in refs)
report={'total_referencias':len(refs),'classificacoes':dict(counts),'ids_novos':[p['id'] for p in plan['papers']],
        'pdfs_movidos':[p['destination'] for p in plan['papers']],'documento_preservado':True,
        'valores_e_estilos_anteriores_preservados':True,'contagens_recalculadas':[24,23,32,2,81],
        'backup':str(OUT/'matriz_artigos_antes_da_inclusao.xlsx')}
(OUT/'resultado_organizacao.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'plano_organizacao.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
